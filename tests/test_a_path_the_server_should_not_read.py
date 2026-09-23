"""A remote server reads and writes only inside the folders it serves.

Two defects in one resolver. Every tool resolved its path with
`Path(s).resolve()`, so any authenticated caller of the deployed server could
name any file in the container -- only a .docx/.xlsx/.pptx extension check
stood in the way. And the caller's string went through `os.path.expandvars`
first: `"$HOME/nonexistent_probe.docx"` came back from the live endpoint as
"File not found: /home/app/nonexistent_probe.docx", so a path naming a
variable that holds a secret would have echoed the secret the same way.

With MCP_CONFINE_PATHS on (the default for every HTTP deployment) a path must
lie inside MCP_OUTPUT_DIR, the workspace root, or MCP_ALLOWED_ROOTS, judged
after symlinks resolve, and is never expanded against the server's
environment. Template and image paths, which bypassed the resolver, now go
through it. A local stdio install is unchanged, and a workspace name is always
a name.
"""

from __future__ import annotations

import os
from pathlib import Path

import pytest

from docx_basic import engine as dbe  # type: ignore[reportMissingImports]
from shared.file_utils import PathOutsideRootError, resolve_path  # type: ignore[reportMissingImports]
from shared.workspace_utils import get_workspace_dir, get_workspace_root  # type: ignore[reportMissingImports]

OUTSIDE = "/etc/probe.docx" if os.name != "nt" else "C:/Windows/probe.docx"


@pytest.fixture
def served(tmp_path, monkeypatch):
    data = tmp_path / "data"
    ws = tmp_path / "ws"
    data.mkdir()
    ws.mkdir()
    monkeypatch.setenv("MCP_CONFINE_PATHS", "1")
    monkeypatch.setenv("MCP_OUTPUT_DIR", str(data))
    monkeypatch.setenv("MCP_WORKSPACE_DIR", str(ws))
    monkeypatch.delenv("MCP_DATA_ROOT", raising=False)
    monkeypatch.delenv("MCP_ALLOWED_ROOTS", raising=False)
    return data


class TestConfined:
    def test_a_file_outside_is_refused(self, served):
        with pytest.raises(PathOutsideRootError, match="outside the folders"):
            resolve_path(OUTSIDE)

    def test_a_relative_path_is_read_from_the_data_folder(self, served):
        assert resolve_path("sub/report.docx") == (served / "sub" / "report.docx").resolve()

    def test_climbing_out_is_refused(self, served):
        with pytest.raises(PathOutsideRootError):
            resolve_path("../escape.docx")

    def test_the_server_environment_is_never_expanded(self, served, monkeypatch):
        monkeypatch.setenv("PROBE_SECRET", "s3cr3t-value")
        path = resolve_path("$PROBE_SECRET/x.docx")
        assert "s3cr3t-value" not in str(path)
        assert path == (served / "$PROBE_SECRET" / "x.docx").resolve()

    def test_a_symlink_is_judged_by_where_it_leads(self, served, tmp_path):
        outside = tmp_path / "outside.docx"
        outside.write_bytes(b"x")
        link = served / "innocent.docx"
        try:
            link.symlink_to(outside)
        except OSError, NotImplementedError:
            pytest.skip("cannot create a symlink here")
        with pytest.raises(PathOutsideRootError):
            resolve_path(str(link))

    def test_an_extra_root_can_be_served(self, served, tmp_path, monkeypatch):
        extra = tmp_path / "extra"
        extra.mkdir()
        monkeypatch.setenv("MCP_ALLOWED_ROOTS", str(extra))
        assert resolve_path(str(extra / "x.docx")) == (extra / "x.docx").resolve()

    def test_a_workspace_base_dir_outside_is_refused(self, served, tmp_path):
        with pytest.raises(PathOutsideRootError, match="base_dir"):
            get_workspace_root(str(tmp_path / "elsewhere"))

    def test_the_tool_refuses(self, served):
        r = dbe.get_document_outline(OUTSIDE)
        assert r["success"] is False
        assert "outside the folders" in r["error"]


class TestLocal:
    def test_a_local_install_is_not_confined(self, monkeypatch):
        monkeypatch.delenv("MCP_CONFINE_PATHS", raising=False)
        assert resolve_path(OUTSIDE) == Path(OUTSIDE).resolve()

    def test_a_local_install_still_expands_variables(self, tmp_path, monkeypatch):
        monkeypatch.delenv("MCP_CONFINE_PATHS", raising=False)
        monkeypatch.setenv("PROBE_DIR", str(tmp_path))
        assert resolve_path("$PROBE_DIR/x.docx") == (tmp_path / "x.docx").resolve()

    @pytest.mark.parametrize("name", ["../../etc", "..", "a/../../b"])
    def test_a_workspace_name_is_always_a_name(self, tmp_path, monkeypatch, name):
        monkeypatch.delenv("MCP_CONFINE_PATHS", raising=False)
        monkeypatch.setenv("MCP_WORKSPACE_DIR", str(tmp_path))
        with pytest.raises(ValueError, match="not a plain name"):
            get_workspace_dir(name)


class TestOutputsConfined:
    """Where a tool WRITES is held to the served folders too.

    A5-sec confined resolve_path, which every INPUT goes through, and missed
    platform_utils.resolve_output_path, the one choke point for all 27 output
    paths: an absolute output_path was used as given and a relative one with a
    folder resolved from the container's working directory. So a confined server
    still created documents wherever a caller pointed it. Found by reading the
    output resolver during a direct sweep of the deployed fleet.
    """

    @pytest.mark.parametrize(
        ("module", "func", "suffix"),
        [
            ("docx_new.engine", "create_document", ".docx"),
            ("xlsx_new.engine", "create_workbook", ".xlsx"),
            ("pptx_new.engine", "create_presentation", ".pptx"),
        ],
    )
    def test_a_creator_refuses_an_outside_output(self, served, tmp_path, module, func, suffix):
        import importlib

        target = tmp_path / "elsewhere" / f"planted{suffix}"
        r = getattr(importlib.import_module(module), func)(str(target), open_after=False)
        assert r["success"] is False, r
        assert "outside the folders" in r["error"]
        assert not target.parent.exists(), "a refused output must leave nothing behind"

    def test_a_relative_output_is_written_into_the_data_folder(self, served):
        from docx_new import engine as dne  # type: ignore[reportMissingImports]

        r = dne.create_document("reports/q3.docx", open_after=False)
        assert r["success"] is True, r
        assert (served / "reports" / "q3.docx").exists()


class TestEveryInputIsResolved:
    """Two tools read their inputs with a bare Path(...).resolve().

    merge_documents(file_paths) and create_from_csv(csv_path) looked a relative
    name up from the process cwd -- so a file the caller had just created in the
    data folder was "not found" -- and took an absolute one from anywhere,
    past the confinement every other input goes through. Found live: merging
    two documents created a moment earlier answered "File not found".
    """

    def test_merge_documents_reads_relative_names_from_the_data_folder(self, served):
        from docx_new import engine as dne  # type: ignore[reportMissingImports]

        assert dne.create_document("a.docx", open_after=False)["success"] is True
        assert dne.create_document("b.docx", open_after=False)["success"] is True
        r = dne.merge_documents(["a.docx", "b.docx"], output_path="ab.docx", open_after=False)
        assert r["success"] is True, r
        assert (served / "ab.docx").exists()

    def test_merge_documents_refuses_an_outside_input(self, served, tmp_path):
        from docx import Document

        from docx_new import engine as dne  # type: ignore[reportMissingImports]

        outside = tmp_path / "private.docx"
        Document().save(str(outside))
        r = dne.merge_documents([str(outside)], output_path="leak.docx", open_after=False)
        assert r["success"] is False, r
        assert "outside the folders" in r["error"]
        assert not (served / "leak.docx").exists()

    def test_create_from_csv_refuses_an_outside_input(self, served, tmp_path):
        from xlsx_new import engine as xne  # type: ignore[reportMissingImports]

        outside = tmp_path / "private.csv"
        outside.write_text("k,v\na,1\n")
        r = xne.create_from_csv(str(outside), output_path="leak.xlsx", open_after=False)
        assert r["success"] is False, r
        assert "outside the folders" in r["error"]
        assert not (served / "leak.xlsx").exists()


class TestImageSourcesGoThroughTheGuard:
    """An image block's URL was fetched with urllib directly.

    No MCP_FETCH_URLS check and no refusal of loopback, private or cloud-metadata
    hosts -- the two things shared/exchange.fetch_url exists to enforce -- so a
    deployed server with URL fetching off still requested any address a caller
    put in an image block. A local image source resolved from the process cwd.
    Found by reading the block renderer during a direct sweep of the fleet.
    """

    @pytest.fixture
    def no_direct_fetch(self, monkeypatch):
        import urllib.request

        def refuse(*a, **k):
            raise AssertionError("fetched a URL directly, past the fleet's guard")

        monkeypatch.setattr(urllib.request, "urlopen", refuse)

    def _blocks(self, source: str) -> dict:
        from docx_new import engine as dne  # type: ignore[reportMissingImports]

        return dne.create_from_blocks(
            title="T", blocks=[{"kind": "image", "path": source}], output_path="img.docx", open_after=False
        )

    def test_a_url_is_refused_when_fetching_is_off(self, served, monkeypatch, no_direct_fetch):
        monkeypatch.delenv("MCP_FETCH_URLS", raising=False)
        r = self._blocks("http://127.0.0.1:9/pixel.png")
        assert "does not fetch URLs" in str(r), r

    def test_a_private_address_is_refused_when_fetching_is_on(self, served, monkeypatch, no_direct_fetch):
        monkeypatch.setenv("MCP_FETCH_URLS", "1")
        monkeypatch.delenv("MCP_FETCH_ALLOW_PRIVATE", raising=False)
        r = self._blocks("http://169.254.169.254/latest/meta-data/x.png")
        assert "non-public address" in str(r), r

    def test_a_local_image_outside_is_refused(self, served, tmp_path):
        outside = tmp_path / "private.png"
        outside.write_bytes(b"\x89PNG\r\n\x1a\n")
        r = self._blocks(str(outside))
        assert "outside the folders" in str(r), r


class TestCreateFromDocxInput:
    def test_a_relative_docx_is_read_from_the_data_folder(self, served):
        from docx_new import engine as dne  # type: ignore[reportMissingImports]
        from pptx_new import engine as pne  # type: ignore[reportMissingImports]

        assert dne.create_from_sections(output_path="src.docx", title="T", sections=[{"heading": "H", "body": "B"}])[
            "success"
        ]
        r = pne.create_from_docx("src.docx", output_path="deck.pptx", open_after=False)
        assert r["success"] is True, r

    def test_an_outside_docx_is_refused(self, served, tmp_path):
        from docx import Document

        from pptx_new import engine as pne  # type: ignore[reportMissingImports]

        outside = tmp_path / "private.docx"
        Document().save(str(outside))
        r = pne.create_from_docx(str(outside), output_path="deck.pptx", open_after=False)
        assert r["success"] is False
        assert "outside the folders" in r["error"]


class TestReceiptsFollowTheDocument:
    """An edit made with a relative path left no receipt.

    append_receipt and read_receipt_log resolved what the caller typed with a
    bare Path(...).resolve(), so a relative name meant the process cwd -- /app
    on a deployed server, where the write failed and was swallowed -- and
    read_receipt then reported an empty log after four edits. Found live.
    """

    def test_an_edit_by_relative_name_is_recorded_and_readable(self, served, monkeypatch, tmp_path):
        from docx_basic import engine as dbe  # type: ignore[reportMissingImports]
        from docx_new import engine as dne  # type: ignore[reportMissingImports]

        elsewhere = tmp_path / "cwd"
        elsewhere.mkdir()
        monkeypatch.chdir(elsewhere)
        assert dne.create_document("doc.docx", open_after=False)["success"]
        assert dbe.append_text("doc.docx", "hello")["success"]
        from docx_basic.helpers import read_receipt_tool  # type: ignore[reportMissingImports]

        log = read_receipt_tool("doc.docx")
        assert log["entries"], log
        assert (served / "doc.docx.mcp_receipt.json").exists()
        assert not list(elsewhere.iterdir())


class TestARefusedPathIsAnAnswer:
    """A refusal raised anywhere inside a tool comes back in the failure shape.

    On the ML server a resolver outside a tool's own `try` let the refusal
    escape: nothing was written, but the caller got "Error executing tool ...:
    Path ... is outside" with no success, op or hint. The per-tool wrapper
    answers it here too, whichever resolver raised it.
    """

    def test_a_refusal_raised_inside_a_tool_is_answered(self):
        from types import SimpleNamespace

        from shared.file_utils import PathOutsideRootError  # type: ignore[reportMissingImports]
        from shared.missing_file import suggest_missing_files  # type: ignore[reportMissingImports]

        def save_it(output_path: str = "") -> dict:
            raise PathOutsideRootError(f"Path '{output_path}' is outside the folders this server can use.")

        registered = SimpleNamespace(name="save_it", fn=save_it)
        mcp = SimpleNamespace(_tool_manager=SimpleNamespace(_tools={"save_it": registered}))
        suggest_missing_files(mcp)
        r = registered.fn(output_path="/elsewhere/out.csv")
        assert r["success"] is False, r
        assert r["op"] == "save_it"
        assert "outside the folders" in r["error"]
        assert "data folder" in r["hint"]

    def test_any_other_error_still_propagates(self):
        from types import SimpleNamespace

        from shared.missing_file import suggest_missing_files  # type: ignore[reportMissingImports]

        def broken() -> dict:
            raise RuntimeError("a real bug")

        registered = SimpleNamespace(name="broken", fn=broken)
        suggest_missing_files(SimpleNamespace(_tool_manager=SimpleNamespace(_tools={"broken": registered})))
        with pytest.raises(RuntimeError, match="a real bug"):
            registered.fn()


class TestAPathOnTheCallersSideIsNamedAsOne:
    """A claude.ai upload path is refused for what it is, with the way in.

    `/mnt/user-data/uploads/Ad_Data.csv` is the only path a chat's model holds
    for an attached file. "Outside the folders this server can use" named the
    rule and sent it guessing folders on a server that cannot see the file.
    """

    def test_the_refusal_says_the_file_is_on_the_callers_side(self, served, monkeypatch):
        monkeypatch.setenv("MCP_FETCH_URLS", "1")
        with pytest.raises(PathOutsideRootError) as caught:
            resolve_path("/mnt/user-data/uploads/Ad_Data.csv")
        message = str(caught.value)
        assert "caller's side" in message
        assert "cannot see it" in message
        assert "link" in message

    def test_any_other_outside_path_keeps_the_plain_refusal(self, served):
        with pytest.raises(PathOutsideRootError, match="outside the folders"):
            resolve_path("/etc/hostname")


class TestAFileSentInlineIsRead:
    """A .docx sent as data:...;base64 is read by a registered tool like any file."""

    def test_the_outline_of_an_inline_document(self, served):
        import base64
        import importlib

        from docx import Document

        doc = Document()
        doc.add_heading("Findings", level=1)
        doc.add_paragraph("Body.")
        built = served.parent / "built.docx"
        doc.save(str(built))
        uri = "data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;name=brief.docx;base64,"
        uri += base64.b64encode(built.read_bytes()).decode()
        tool = importlib.import_module("servers.docx_basic.docx_basic.server").mcp._tool_manager._tools
        r = tool["get_document_outline"].fn(file_path=uri)
        assert r["success"] is True, r
        assert "Findings" in str(r)
        assert "base64" not in str(r)
        assert (served / "inbox" / "brief.docx").exists()
