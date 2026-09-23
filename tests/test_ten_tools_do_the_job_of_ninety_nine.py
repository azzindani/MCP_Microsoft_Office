"""Every Office tool as ten domain tools on one endpoint, `action` plus `args`.

A model connected to all eleven tiers reads 99 tool names on every turn, nine
of them twice for different file types. The domain endpoint lists ten -- read,
edit, create for each of Word, Excel and PowerPoint, and one for history --
and each `action` is a tier tool by its own name, run by the tier's own
`run()`, so validation, wrappers and answers are the tier's.
"""

from __future__ import annotations

import asyncio
import collections
import importlib

import pytest

from servers.office_domain.office_domain.server import DOMAINS, mcp

TIERS = [
    "docx_basic",
    "docx_tables",
    "docx_layout",
    "docx_new",
    "xlsx_basic",
    "xlsx_formulas",
    "xlsx_charts",
    "xlsx_new",
    "pptx_basic",
    "pptx_design",
    "pptx_new",
]


def _tier_listed() -> collections.Counter:
    names: collections.Counter = collections.Counter()
    for tier in TIERS:
        server = importlib.import_module(f"servers.{tier}.{tier}.server").mcp
        names.update(t.name for t in asyncio.run(server.list_tools()))
    return names


def _listed():
    return {t.name: t for t in asyncio.run(mcp.list_tools())}


def _call(tool: str, action: str, args: dict) -> dict:
    return asyncio.run(mcp._tool_manager._tools[tool].run({"action": action, "args": args}))


@pytest.fixture
def files(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    doc = Document()
    doc.add_heading("Findings", level=1)
    doc.add_paragraph("Body.")
    docx = tmp_path / "brief.docx"
    doc.save(str(docx))
    book = Workbook()
    book.active.title = "Q3"
    book.active["A1"] = "units"
    xlsx = tmp_path / "q3.xlsx"
    book.save(str(xlsx))
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Hello"
    pptx = tmp_path / "deck.pptx"
    deck.save(str(pptx))
    return {"docx": str(docx), "xlsx": str(xlsx), "pptx": str(pptx)}


class TestTenNotNinetyNine:
    def test_ten_tools_are_listed(self):
        assert sorted(_listed()) == sorted(DOMAINS) and len(_listed()) == 10

    def test_the_actions_are_exactly_the_tier_tools(self):
        actions = collections.Counter(
            a for t in _listed().values() for a in t.inputSchema["properties"]["action"]["enum"]
        )
        assert actions == _tier_listed(), "a tier tool is missing from the domains, or one is in twice"

    def test_a_name_shared_by_file_types_is_the_right_tool_in_each(self):
        docx = mcp._tool_manager._tools["docx_edit"]
        xlsx = mcp._tool_manager._tools["xlsx_edit"]
        assert "set_cell" in docx.parameters["properties"]["action"]["enum"]
        assert "set_cell" in xlsx.parameters["properties"]["action"]["enum"]
        docx_tier = importlib.import_module("servers.docx_tables.docx_tables.server").mcp._tool_manager._tools
        xlsx_tier = importlib.import_module("servers.xlsx_basic.xlsx_basic.server").mcp._tool_manager._tools
        assert docx_tier["set_cell"].parameters != xlsx_tier["set_cell"].parameters

    def test_the_schema_is_the_pipeline_shape(self):
        for tool in _listed().values():
            assert tool.inputSchema["required"] == ["action"]
            assert tool.inputSchema["properties"]["args"]["additionalProperties"] is False
            for action in tool.inputSchema["properties"]["action"]["enum"]:
                assert f"- {action}:" in tool.description, (tool.name, action)


class TestAnActionIsTheTierTool:
    @pytest.mark.parametrize(
        ("tool", "action", "kind"),
        [
            ("docx_read", "get_document_outline", "docx"),
            ("xlsx_read", "list_sheets", "xlsx"),
            ("pptx_read", "read_presentation", "pptx"),
            ("office_history", "get_history", "docx"),
        ],
    )
    def test_it_answers_as_the_tier_does(self, files, tool, action, kind):
        result = _call(tool, action, {"file_path": files[kind]})
        assert result["success"] is True, result


class TestARefusalSaysWhatToDo:
    def test_a_shared_name_asked_of_the_wrong_tool_names_both_homes(self):
        result = _call("docx_read", "set_cell", {})
        assert result["success"] is False
        assert "docx_edit" in result["hint"] and "xlsx_edit" in result["hint"]

    def test_an_unknown_action_lists_the_actions(self):
        result = _call("pptx_read", "no_such_thing", {})
        assert result["success"] is False and "read_slide" in result["hint"]

    def test_an_argument_the_action_does_not_take_is_named(self, files):
        result = _call("docx_read", "get_document_outline", {"file_path": files["docx"], "colour": "red"})
        assert result["success"] is False and "colour" in result["error"]

    def test_a_missing_required_argument_is_an_answer(self):
        result = _call("xlsx_read", "list_sheets", {})
        assert result["success"] is False and "file_path" in result["hint"]


def test_the_unified_server_serves_it_at_mcp_and_a_connector_can_sign_in(tmp_path):
    """Run in a fresh interpreter: the OAuth bridge exists only when a key is set at import."""
    import json
    import os
    import subprocess
    import sys
    from pathlib import Path

    root = Path(__file__).resolve().parents[1]
    env = {**os.environ, "OFFICE_API_KEY": "test-key-not-a-secret", "OFFICE_PUBLIC_URL": "https://office.example.test"}
    for tier in [*TIERS, "domain"]:
        env[f"OFFICE_{tier.upper()}_OAUTH_STATE_DIR"] = str(tmp_path / tier)
    probe = (
        "import json\n"
        "from starlette.testclient import TestClient\n"
        "import unified_server\n"
        "with TestClient(unified_server.app) as c:\n"
        "    out = {p: c.get(p).json() for p in ('/', '/.well-known/oauth-protected-resource/mcp',\n"
        "        '/.well-known/oauth-protected-resource', '/.well-known/oauth-authorization-server',\n"
        "        '/docx-basic/.well-known/oauth-protected-resource')}\n"
        "    out['status'] = c.post('/mcp', json={}, headers={'Accept': 'application/json, text/event-stream'}).status_code\n"
        "print(json.dumps(out))\n"
    )
    done = subprocess.run([sys.executable, "-c", probe], cwd=root, env=env, capture_output=True, text=True, timeout=120)
    assert done.returncode == 0, done.stderr[-2000:]
    out = json.loads(done.stdout.strip().splitlines()[-1])
    assert out["/"]["mcp"] == "/mcp"
    at_mcp = out["/.well-known/oauth-protected-resource/mcp"]
    assert at_mcp["resource"].endswith("/mcp")
    assert out["/.well-known/oauth-protected-resource"] == at_mcp
    assert at_mcp["authorization_servers"] == [out["/.well-known/oauth-authorization-server"]["issuer"]]
    assert out["status"] == 401
    assert out["/docx-basic/.well-known/oauth-protected-resource"]["resource"].endswith("/docx-basic/mcp")
