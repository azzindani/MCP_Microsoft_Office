"""A call that trusts the schema and the words works, and lands where it should.

Sweep findings on this server, each a call that followed what the tool said:

- create_agenda(items=["Budget", "Hiring"]) -- a list of topics, the natural
  input -- answered "'str' object has no attribute 'get'".
- set_conditional_format's description said rule gt/lt/between/eq while the
  schema said greater_than/...; `color` was a closed set tools/list never named.
- add_chart: tools/list marked anchor_cell optional and the runtime refused
  without it.
- update_chart / delete_chart: an out-of-range chart_index was answered "Use
  add_chart to see chart indices" -- add_chart makes a chart and lists nothing.
- add_text_box's default spot sat over a slide's title, and a second call
  landed exactly on the first.
"""

from __future__ import annotations

import openpyxl
import pytest
from pptx import Presentation

from servers.pptx_basic.pptx_basic import engine as pptx_basic
from servers.pptx_new.pptx_new import engine as pptx_new
from servers.xlsx_charts.xlsx_charts import engine as xlsx_charts
from servers.xlsx_formulas.xlsx_formulas import server as xlsx_formulas_server


@pytest.fixture(autouse=True)
def _output(tmp_path, monkeypatch):
    monkeypatch.setenv("MCP_OUTPUT_DIR", str(tmp_path))


class TestAnAgendaTakesTopics:
    def test_a_list_of_strings(self, tmp_path):
        out = tmp_path / "agenda.pptx"
        result = pptx_new.create_agenda(str(out), "Weekly", "2026-09-23", ["Budget", "Hiring"], open_after=False)
        assert result["success"] is True, result.get("error")
        body = Presentation(str(out)).slides[1].placeholders[1].text_frame.text
        assert "Budget" in body and "Hiring" in body and "()" not in body

    def test_a_dict_without_a_topic_is_refused_before_writing(self, tmp_path):
        out = tmp_path / "agenda.pptx"
        result = pptx_new.create_agenda(str(out), "Weekly", "2026-09-23", [{"owner": "Ana"}], open_after=False)
        assert result["success"] is False and "topic" in result["error"] and not out.exists()

    def test_the_full_form_still_reads_every_part(self, tmp_path):
        out = tmp_path / "agenda.pptx"
        items = [{"topic": "Budget", "duration": "10 min", "owner": "Ana"}]
        pptx_new.create_agenda(str(out), "Weekly", "2026-09-23", items, open_after=False)
        body = Presentation(str(out)).slides[1].placeholders[1].text_frame.text
        assert "Budget (10 min) — Ana" in body


class TestConditionalFormatSaysItsVocabulary:
    def _tool(self):
        return xlsx_formulas_server.mcp._tool_manager._tools["set_conditional_format"]

    def test_color_is_a_published_enum(self):
        assert self._tool().parameters["properties"]["color"]["enum"] == ["blue", "green", "red", "yellow"]

    def test_the_description_uses_the_schemas_rule_names(self):
        text = self._tool().description
        assert "greater_than" in text and "less_than" in text


@pytest.fixture
def workbook(tmp_path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Region", "Units"], ["North", 3], ["South", 5], ["East", 2]]:
        ws.append(row)
    path = tmp_path / "sales.xlsx"
    wb.save(path)
    return str(path)


class TestAChartWithoutAnAnchor:
    def test_it_goes_beside_the_data(self, workbook):
        result = xlsx_charts.add_chart(workbook, "Data", "bar", "A1:B4", title="Units")
        assert result["success"] is True, result.get("error")
        assert result["anchor_cell"] == "D1" and "anchor_note" in result

    def test_a_given_anchor_is_kept(self, workbook):
        result = xlsx_charts.add_chart(workbook, "Data", "bar", "A1:B4", anchor_cell="F10")
        assert result["anchor_cell"] == "F10" and "anchor_note" not in result


class TestAWrongChartIndexListsTheCharts:
    @pytest.mark.parametrize("call", [xlsx_charts.delete_chart, xlsx_charts.update_chart])
    def test_an_empty_sheet_says_so(self, workbook, call):
        result = call(workbook, "Data", 0)
        assert result["success"] is False
        assert "has no charts" in result["hint"] and "see chart indices" not in result["hint"]

    @pytest.mark.parametrize("call", [xlsx_charts.delete_chart, xlsx_charts.update_chart])
    def test_the_charts_are_named_by_index(self, workbook, call):
        xlsx_charts.add_chart(workbook, "Data", "bar", "A1:B4", title="Units", anchor_cell="D1")
        result = call(workbook, "Data", 5)
        assert "1 chart(s): 0 'Units'" in result["hint"]


@pytest.fixture
def deck(tmp_path) -> str:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "Quarterly results"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return str(path)


class TestATextBoxLandsClear:
    def test_the_default_spot_moves_off_the_title(self, deck):
        result = pptx_basic.add_text_box(deck, 0, "A note")
        assert result["success"] is True, result.get("error")
        assert result["position"]["top"] > 1.0 and "placement_note" in result

    def test_a_second_default_box_does_not_land_on_the_first(self, deck):
        # It used to land exactly on it. Placement keeps shapes clear of what
        # they DRAW -- a one-line box draws well short of its 1in frame -- so
        # the second goes below the first one's text, not below its frame.
        first = pptx_basic.add_text_box(deck, 0, "First")["position"]
        second = pptx_basic.add_text_box(deck, 0, "Second")["position"]
        assert second["top"] > first["top"] + 0.3

    def test_a_chosen_position_is_kept_and_the_overlap_said(self, deck):
        result = pptx_basic.add_text_box(deck, 0, "Caption", left=0.5, top=0.5, width=9.0, height=1.5)
        assert result["position"]["top"] == 0.5
        assert "over 1 existing shape" in result["placement_note"]
