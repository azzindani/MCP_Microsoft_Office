"""The response said Heading 1. The document said Normal.

Round 28 sent every dispatch parameter in the fleet a value it cannot mean.
`append_text` and `insert_paragraph` were two of the five that answered anyway:

    append_text(style="Heading 1")  -> success, style: "Heading 1"
    append_text(style="Headng 1")   -> success, style: "Headng 1"

    read_document:
      index 1  "should be a heading"  style: "Heading 1"   <- applied
      index 2  "typo style"           style: "Normal"      <- silently fell back

Both tools did this:

    try:
        para.style = doc.styles[style]
    except KeyError:
        pass  # Default style

and then reported the caller's own spelling back under `success: True`. One
typo in a style name produced body text where a heading was asked for, and
every downstream tool then agreed with the document rather than the response --
`get_document_outline` will not list it, `get_document_index` opens no section
for it.

The style is now checked against the open document before the snapshot is taken
and before anything is written, so a refusal costs nothing and the file is
untouched.

The other three fixes here are the same shape: a response that described the
request rather than the result.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from servers.docx_basic.docx_basic import engine as basic  # noqa: E402
from servers.docx_new.docx_new import engine as new  # noqa: E402
from servers.docx_tables.docx_tables import engine as tables  # noqa: E402


@pytest.fixture
def doc(tmp_path):
    path = tmp_path / "d.docx"
    r = new.create_from_text(str(path), [{"text": "seed", "style": "Normal"}], open_after=False)
    assert r["success"] is True, r.get("error")
    return str(path)


class TestAStyleIsAppliedOrRefused:
    def test_a_real_style_is_applied(self, doc):
        r = basic.append_text(doc, "a heading", style="Heading 1")
        assert r["success"] is True, r.get("error")
        read = basic.read_document(doc)
        assert read["paragraphs"][-1]["style"] == "Heading 1"

    def test_a_typo_is_refused(self, doc):
        r = basic.append_text(doc, "not a heading", style="Headng 1")
        assert r["success"] is False, "a style the document does not define was reported as applied"

    def test_the_refusal_suggests_the_real_name(self, doc):
        r = basic.append_text(doc, "x", style="Headng 1")
        assert "Heading 1" in r["hint"], r["hint"]

    def test_nothing_is_written_when_the_style_is_refused(self, doc):
        before = basic.read_document(doc)["paragraph_count"]
        basic.append_text(doc, "x", style="Headng 1")
        assert basic.read_document(doc)["paragraph_count"] == before

    def test_insert_paragraph_behaves_the_same(self, doc):
        r = basic.insert_paragraph(doc, 0, "x", style="Headng 1")
        assert r["success"] is False
        assert "Heading 1" in r["hint"], r["hint"]

    def test_insert_paragraph_still_inserts_with_a_real_style(self, doc):
        r = basic.insert_paragraph(doc, 0, "inserted", style="Heading 2")
        assert r["success"] is True, r.get("error")
        assert basic.read_document(doc)["paragraphs"][1]["style"] == "Heading 2"


class TestAddTableSaysWhichTableItMade:
    """Every other tool here addresses tables by index; this one returned none."""

    def test_it_returns_the_index(self, doc):
        r = tables.add_table(doc, after_paragraph_index=-1, rows=2, cols=2, data=[["a", "b"], ["c", "d"]])
        assert r["success"] is True, r.get("error")
        assert r["table_index"] == 0
        assert r["table_count"] == 1

    def test_the_index_it_returns_is_the_table_it_made(self, doc):
        """Two tables built the same way are identical by value; only the index tells them apart."""
        tables.add_table(doc, after_paragraph_index=0, rows=1, cols=1, data=[["first"]])
        second = tables.add_table(doc, after_paragraph_index=-1, rows=1, cols=1, data=[["second"]])
        read = tables.read_table(doc, second["table_index"])
        assert read["data"][0][0]["text"] == "second", read["data"]

    def test_it_warns_when_existing_tables_are_renumbered(self, doc):
        """anchor=-1 goes in front of a table anchored later, so held indices move.

        This is the sequence that cost round 28 a row out of the wrong table:
        holding table_index=1 across an add_table and then calling delete_row.
        """
        first = tables.add_table(doc, after_paragraph_index=0, rows=1, cols=1, data=[["first"]])
        assert first["table_index"] == 0
        r = tables.add_table(doc, after_paragraph_index=-1, rows=1, cols=1, data=[["second"]])
        assert r["table_index"] == 0, "the new table took index 0"
        warned = " ".join(str(p) for p in r["progress"])
        assert "renumber" in warned.lower(), warned

    def test_the_old_index_really_does_move(self, doc):
        """The warning is only worth having if it describes something true."""
        tables.add_table(doc, after_paragraph_index=0, rows=1, cols=1, data=[["first"]])
        tables.add_table(doc, after_paragraph_index=-1, rows=1, cols=1, data=[["second"]])
        assert tables.read_table(doc, 1)["data"][0][0]["text"] == "first"


class TestTheInvoiceSaysWhatIsInTheFile:
    """Line total, subtotal, tax and grand total are all uncalculated formulas."""

    def test_it_reports_tax_and_total_not_just_subtotal(self, tmp_path):
        from servers.xlsx_new.xlsx_new import engine as xlsx

        r = xlsx.create_invoice(
            output_path=str(tmp_path / "inv.xlsx"),
            company_name="Analytics Ltd",
            client_name="Marketing",
            invoice_number="R28-001",
            items=[{"description": "Review", "quantity": 2, "unit_price": 500}],
            tax_rate=0.1,
            open_after=False,
        )
        assert r["success"] is True, r.get("error")
        assert r["subtotal"] == 1000.0
        assert r["tax"] == 100.0
        assert r["total"] == 1100.0

    def test_it_says_the_cells_hold_formulas_with_no_value_yet(self, tmp_path):
        from servers.xlsx_new.xlsx_new import engine as xlsx

        r = xlsx.create_invoice(
            output_path=str(tmp_path / "inv.xlsx"),
            company_name="A",
            client_name="B",
            invoice_number="1",
            items=[{"description": "x", "quantity": 1, "unit_price": 1}],
            tax_rate=0.1,
            open_after=False,
        )
        assert r["calculated"] is False
        assert "no calculation engine" in r["note"]

    def test_the_claim_matches_the_file(self, tmp_path):
        """The note is only worth having if the cells really are uncalculated."""
        import openpyxl

        from servers.xlsx_new.xlsx_new import engine as xlsx

        out = tmp_path / "inv.xlsx"
        xlsx.create_invoice(
            output_path=str(out),
            company_name="A",
            client_name="B",
            invoice_number="1",
            items=[{"description": "x", "quantity": 1, "unit_price": 1}],
            tax_rate=0.1,
            open_after=False,
        )
        ws = openpyxl.load_workbook(str(out)).active
        formulas = [
            c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")
        ]
        assert formulas, "the note claims formulas; the file has none"


class TestOneNameForTheFillColour:
    """docx_tables spelled it `fill`, xlsx_charts `fill_color`, same tool name."""

    def test_the_docx_styler_takes_the_spreadsheet_spelling(self, doc):
        tables.add_table(doc, after_paragraph_index=-1, rows=2, cols=2, data=[["a", "b"], ["c", "d"]])
        r = tables.set_cell_style(doc, 0, fill_color="D9E2F3", row=0, col=-1)
        assert r["success"] is True, r.get("error")

    def test_the_spreadsheet_styler_takes_the_document_spelling(self, tmp_path):
        from servers.xlsx_charts.xlsx_charts import engine as charts
        from servers.xlsx_new.xlsx_new import engine as xlsx

        out = tmp_path / "b.xlsx"
        xlsx.create_from_data(str(out), "Data", ["a"], [[1]], open_after=False)
        r = charts.set_cell_style(str(out), "Data", "A1", fill="FFFF00", open_after=False)
        assert r["success"] is True, r.get("error")
